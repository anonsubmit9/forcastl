"""CSV and manifest output functions for detection runs."""

import csv
import json
import re
from pathlib import Path
from typing import IO, Iterable, List, Dict, Optional
from datetime import datetime
from colorama import Fore, Style

from forcastl import config


# Column glossary embedded at the top of every detection_results_*.csv. Lines
# are written with a leading '#' so pandas readers can skip via comment='#'
# and our own readers (see read_detection_csv below) skip them automatically.
_DETECTION_CSV_GLOSSARY = [
    "DETECTION BENCHMARK RESULTS - column glossary",
    "----------------------------------------------------------------------",
    "Model               LLM model id used to score this run",
    "Filename            Source artifact analyzed (EVTX file or CSV stem)",
    "Difficulty          Tag from metadata.json: easy / medium / hard",
    "Test Set            A=Known Pattern, B=Obfuscated, C=Noise-heavy, D=Benign",
    "Status              processed | error | context_exceeded",
    "Malicious           Model verdict: YES (attack) or NO (routine/benign)",
    "Confidence %        Heuristic confidence from response structure quality, not model-stated (0-100)",
    "Alignment %         Share of cited evidence found verbatim in source (0-100)",
    "Structure Valid     Response parsed cleanly into the required sections",
    "Event IDs           Event IDs extracted from the model's EVIDENCE block",
    "Processes           Process names extracted from the model's EVIDENCE block",
    "Accounts            Account names extracted from the model's EVIDENCE block",
    "Commands            Command lines / script blocks extracted from EVIDENCE",
    "Network             IPs / hostnames / shares extracted from EVIDENCE",
    "Registry            Registry keys / values extracted from EVIDENCE",
    "Hallucination %     Share of cited evidence not found in source events (0-100)",
    "Response Time (s)   Wall-clock seconds from prompt sent to response received",
    "Extraction (6)      Did the model find the right evidence? (0-6)",
    "Interpretation (6)  Did it correctly explain what the evidence implies? (0-6)",
    "NoHallucination (4) 4 = no fabricated evidence, 0 = heavy fabrication",
    "Reasoning (4)       Quality of the chain of reasoning in EXPLANATION (0-4)",
    "Total (20)          Sum of the four criteria above (0-20)",
    "Grade               Letter grade from Total: A>=17, B>=14, C>=10, D<10",
    "LLM Response        Full raw text the model returned, before any parsing",
    "",
    "Skip these lines in pandas:  pd.read_csv(path, comment='#')",
    "----------------------------------------------------------------------",
]


def _write_glossary(f: IO[str]) -> None:
    for line in _DETECTION_CSV_GLOSSARY:
        f.write(f"# {line}\n" if line else "#\n")


# ── User-facing (download from GUI) view ──
# The on-disk CSV keeps every column for the comparison report's benefit.
# When the GUI hands an artifact to a human for troubleshooting, we ship a
# trimmed XLSX: drop the per-file scoring columns (those metrics are
# aggregated into the GUI's verdict + KPIs, not reviewed per-file) and put
# the standard request on its own sheet so the reader knows what prompt
# produced the `LLM Response` column.

_USER_FACING_COLUMNS = [
    'Model', 'Filename', 'Difficulty', 'Test Set',
    'Status', 'Malicious', 'Structure Valid',
    'Event IDs', 'Processes', 'Accounts',
    'Commands', 'Network', 'Registry',
    'Hallucination %',
    'Response Time (s)', 'LLM Response',
]


def _build_request_template_lines() -> List[str]:
    """The static prompt sent to every model on every file in this run.

    Built once from PromptManager so the documentation can never drift from
    the actual prompt. The variable XML payload is replaced with a marker.
    """
    from forcastl.core.prompt import PromptManager
    template = PromptManager().generate_simple_malicious_prompt(
        "[per-file event-log data inserted here]"
    )
    return template.splitlines()


def _backfill_evidence_columns(rows: List[Dict[str, str]]) -> List[Dict[str, str]]:
    """Recover Commands / Network / Registry for rows whose source CSV was
    written before those columns existed.

    The detection_results CSV has always carried the full ``LLM Response``
    text, so we can re-run the response parser on each row and pull the
    structured EVIDENCE block back out. Hallucination % can't be recovered
    this way (it depends on the original event-extraction state, which the
    CSV doesn't carry) — for those rows the column stays blank.

    No-op for rows that already have the columns populated, so it's safe to
    call unconditionally.
    """
    backfill_keys = ("Commands", "Network", "Registry")
    needs_any = any(
        any(not r.get(k) for k in backfill_keys) and r.get("LLM Response")
        for r in rows
    )
    if not needs_any:
        return rows

    from forcastl.core.response_parser import parse_structured_response
    enriched: List[Dict[str, str]] = []
    for r in rows:
        out = dict(r)
        if r.get("LLM Response") and any(not out.get(k) for k in backfill_keys):
            try:
                parsed = parse_structured_response(r["LLM Response"])
                ev = parsed.get("evidence") or {}
                if not out.get("Commands"):
                    out["Commands"] = " | ".join(ev.get("commands") or [])
                if not out.get("Network"):
                    out["Network"] = ", ".join(ev.get("network") or [])
                if not out.get("Registry"):
                    out["Registry"] = " | ".join(ev.get("registry") or [])
            except Exception:
                pass  # leave row untouched on any parser hiccup
        enriched.append(out)
    return enriched


def render_user_facing_xlsx(rows: List[Dict[str, str]],
                            run_meta: Optional[Dict] = None) -> bytes:
    """Build a 2-sheet workbook: 'Request' (the prompt template) and 'Results'
    (the trimmed data table). Same data as render_user_facing_csv, but split
    so the prompt isn't crammed at the top of the data sheet.
    """
    import io
    rows = _backfill_evidence_columns(rows)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as e:
        raise RuntimeError(
            "openpyxl not installed. Add `openpyxl>=3.1.0` to requirements.txt "
            "and run `pip install -r requirements.txt`."
        ) from e

    wb = Workbook()

    # ── Sheet 1: Request — the standard prompt template ──
    req_ws = wb.active
    req_ws.title = "Request"
    bold = Font(bold=True)
    title_font = Font(bold=True, size=14)
    section_font = Font(bold=True, size=11, color="3B5CCC")
    muted_font = Font(color="595959", italic=True)

    req_ws["A1"] = "Standard request"
    req_ws["A1"].font = title_font
    if run_meta:
        bits = []
        if run_meta.get("model"):
            bits.append(f"Model: {run_meta['model']}")
        if run_meta.get("run_id"):
            bits.append(f"Run: {run_meta['run_id']}")
        if bits:
            req_ws["A2"] = " · ".join(bits)
            req_ws["A2"].font = muted_font
    req_ws["A3"] = ("This is the exact prompt sent to the model for every file in this run. "
                    "Variable parts are marked with [bracketed placeholders]; the event-log "
                    "data is substituted in per file.")
    req_ws["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    req_ws.row_dimensions[3].height = 45
    req_ws["A5"] = "Prompt"
    req_ws["A5"].font = section_font

    template_text = "\n".join(_build_request_template_lines())
    req_ws["A6"] = template_text
    req_ws["A6"].alignment = Alignment(wrap_text=True, vertical="top")
    # Tall enough that the entire prompt is visible on screen without resizing.
    req_ws.row_dimensions[6].height = max(400, 14 * (template_text.count("\n") + 1))
    req_ws.column_dimensions["A"].width = 110

    # ── Sheet 2: Results — the data table ──
    data_ws = wb.create_sheet(title="Results")
    header_fill = PatternFill(start_color="F4F4F2", end_color="F4F4F2", fill_type="solid")
    for col_idx, header in enumerate(_USER_FACING_COLUMNS, start=1):
        cell = data_ws.cell(row=1, column=col_idx, value=header)
        cell.font = bold
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left", vertical="center")
    for row_idx, r in enumerate(rows, start=2):
        for col_idx, header in enumerate(_USER_FACING_COLUMNS, start=1):
            value = r.get(header, "")
            cell = data_ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="top",
                                       wrap_text=(header == "LLM Response"))
    # Reasonable column widths — keep filename + LLM response visible.
    widths = {
        "Model": 20, "Filename": 50, "Difficulty": 12, "Test Set": 10,
        "Status": 12, "Malicious": 11, "Structure Valid": 14,
        "Event IDs": 18, "Processes": 22, "Accounts": 22,
        "Commands": 28, "Network": 22, "Registry": 28,
        "Hallucination %": 14,
        "Response Time (s)": 16, "LLM Response": 80,
    }
    for col_idx, header in enumerate(_USER_FACING_COLUMNS, start=1):
        letter = get_column_letter(col_idx)
        data_ws.column_dimensions[letter].width = widths.get(header, 16)
    data_ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_detection_csv(path) -> List[Dict[str, str]]:
    """Read a detection_results_*.csv, transparently skipping the '#' glossary
    block at the top. Use this in place of csv.DictReader on this file family
    so callers don't have to know about the comment header.
    """
    with open(path, "r", newline="", encoding="utf-8") as f:
        return list(csv.DictReader(_iter_data_lines(f)))


def _iter_data_lines(f: Iterable[str]) -> Iterable[str]:
    for line in f:
        if not line.startswith("#"):
            yield line
            break
    yield from f


def generate_csv(results: List[Dict], *, model_name: str, run_id: str = None,
                 lookup_file_tag=None) -> str:
    """Generate CSV export"""
    timestamp = run_id or datetime.now().strftime('%Y%m%d_%H%M%S')
    safe_model = re.sub(r'[^\w\-.]', '_', model_name)
    config.OUTPUTS_DIR.mkdir(parents=True, exist_ok=True)
    csv_file = str(config.OUTPUTS_DIR / f"detection_results_{safe_model}_{timestamp}.csv")

    with open(csv_file, 'w', newline='', encoding='utf-8') as f:
        _write_glossary(f)
        writer = csv.writer(f)
        writer.writerow([
            'Model', 'Filename', 'Difficulty', 'Test Set',
            'Status',
            'Malicious', 'Confidence %', 'Alignment %',
            'Structure Valid',
            'Event IDs', 'Processes', 'Accounts',
            'Commands', 'Network', 'Registry',
            'Hallucination %',
            'Response Time (s)',
            'Extraction (6)', 'Interpretation (6)', 'NoHallucination (4)',
            'Reasoning (4)', 'Total (20)', 'Grade',
            'LLM Response'
        ])

        for result in results:
            detection = result.get('detection', {})
            parsed = detection.get('parsed', {})
            validation = detection.get('validation', {})
            alignment = detection.get('alignment', {})
            score_bd = detection.get('score_breakdown', {})
            hallucination = detection.get('hallucination') or {}

            filename = Path(result['file_path']).name
            difficulty = lookup_file_tag(filename, 'difficulty') or '' if lookup_file_tag else ''
            test_set_tag = lookup_file_tag(filename, 'test_set') or '' if lookup_file_tag else ''
            status = result.get('status', 'processed') or 'processed'
            is_malicious = 'YES' if detection.get('is_malicious') else 'NO'
            confidence = detection.get('confidence', 0)
            align_score = alignment.get('score') if alignment else 0
            if align_score is None:
                align_score = 0
            structure_valid = 'YES' if validation.get('is_valid') else 'NO'

            evidence = parsed.get('evidence', {}) or {}
            event_ids = ', '.join(evidence.get('event_ids', []) or [])
            processes = ', '.join(evidence.get('processes', []) or [])
            accounts = ', '.join(evidence.get('accounts', []) or [])
            commands = ' | '.join(evidence.get('commands', []) or [])
            network = ', '.join(evidence.get('network', []) or [])
            registry = ' | '.join(evidence.get('registry', []) or [])

            # Per-file hallucination percentage. `ratio` is 0..1 of evidence
            # claims that weren't grounded in the actual events. Empty string
            # for unprocessed/errored files (no detection ran).
            hall_ratio = hallucination.get('ratio')
            hall_pct = (f"{hall_ratio * 100:.1f}"
                        if isinstance(hall_ratio, (int, float)) else '')

            response_time = result.get('response_time', 0)
            llm_response = result.get('llm_response', '')

            writer.writerow([
                model_name, filename, difficulty, test_set_tag,
                status,
                is_malicious, f"{confidence:.1f}", f"{align_score:.1f}",
                structure_valid,
                event_ids, processes, accounts,
                commands, network, registry,
                hall_pct,
                f"{response_time:.2f}",
                score_bd.get('extraction', ''), score_bd.get('interpretation', ''),
                score_bd.get('hallucination', ''), score_bd.get('reasoning', ''),
                score_bd.get('total', ''), score_bd.get('grade', ''),
                llm_response
            ])

    return csv_file


def generate_scoring_csv(all_runs: List[List[Dict]], *, model_name: str,
                         lookup_file_tag=None):
    """Generate multi-run scoring sheet CSV."""
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    safe_model = re.sub(r'[^\w\-.]', '_', model_name)
    config.OUTPUTS_DIR.mkdir(parents=True, exist_ok=True)
    csv_file = str(config.OUTPUTS_DIR / f"scoring_sheet_{safe_model}_{timestamp}.csv")

    with open(csv_file, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow([
            'TestID', 'Difficulty', 'Model', 'Run', 'Status', 'Extraction', 'Interpretation',
            'NoHallucination', 'Reasoning', 'TotalScore', 'Grade'
        ])

        for run_num, results in enumerate(all_runs, 1):
            for result in results:
                status = result.get('status', 'processed') or 'processed'
                detection = result.get('detection', {})
                score_bd = detection.get('score_breakdown', {})
                filename = Path(result['file_path']).name
                difficulty = lookup_file_tag(filename, 'difficulty') or '' if lookup_file_tag else ''

                writer.writerow([
                    filename, difficulty, model_name, run_num, status,
                    score_bd.get('extraction', ''),
                    score_bd.get('interpretation', ''),
                    score_bd.get('hallucination', ''),
                    score_bd.get('reasoning', ''),
                    score_bd.get('total', ''),
                    score_bd.get('grade', '')
                ])

    print(f"{Fore.GREEN}[OK] Scoring Sheet: {csv_file}{Style.RESET_ALL}")


def write_run_manifest(
    *,
    run_id: str,
    started_at: str,
    finished_at: str,
    mode: Optional[str],
    sample_size: Optional[int],
    results: List[Dict],
    outputs: Dict[str, str],
    model_id: str,
    input_format: str,
    csv_dir: Optional[str],
    max_tokens: int,
    timeout: int,
    delay: float,
    difficulty: Optional[str],
    test_set: Optional[str],
    server_url: Optional[str],
    selected_files: Optional[List[Dict]] = None,
    recommendations: Optional[List[Dict]] = None,
    compute_run_summary=None,
) -> str:
    """Write a minimal machine-readable manifest for the run."""
    summary = compute_run_summary(results)

    files = []
    for r in results:
        filename = Path(r.get("file_path", "")).name
        det = r.get("detection", {}) or {}
        score_bd = det.get("score_breakdown", {}) or {}
        align = det.get("alignment", {}) or {}
        hall = det.get("hallucination", {}) or {}
        files.append({
            "filename": filename,
            "status": (r.get("status") or "processed"),
            "malicious": bool(det.get("is_malicious")) if r.get("status") == "processed" else None,
            "confidence_pct": det.get("confidence") if r.get("status") == "processed" else None,
            "alignment_pct": align.get("score") if r.get("status") == "processed" else None,
            "score_20": score_bd.get("total") if r.get("status") == "processed" else None,
            "grade": score_bd.get("grade") if r.get("status") == "processed" else None,
            "hallucination_rate": hall.get("ratio") if r.get("status") == "processed" else None,
            "hallucinated_fields": hall.get("total_hallucinated") if r.get("status") == "processed" else None,
            "claimed_fields": hall.get("total_claimed") if r.get("status") == "processed" else None,
            "response_time_s": r.get("response_time", 0),
        })

    selection = None
    if selected_files is not None:
        selection = [Path(fi.get("full_path", "")).name for fi in selected_files if fi]

    from forcastl.benchmark_contract import SCORING_VERSION

    manifest = {
        "schema_version": 1,
        # Stamp the scoring rubric version so runs scored under different rubrics are distinguishable
        # (they are NOT score-comparable). Detection manifests written before
        # 2026-06-09 lack this key — treat its absence as v1.
        "scoring_version": SCORING_VERSION,
        "run_id": run_id,
        "kind": "detection",
        "raw_only": True,
        "server": server_url,
        "model": model_id,
        "input_format": input_format,
        "csv_dir": csv_dir,
        "mode": mode,
        "sample_size": sample_size,
        "max_output_tokens": max_tokens,
        "timeout_s": timeout,
        "delay_s": delay,
        "difficulty_filter": difficulty,
        "test_set_filter": test_set,
        "started_at": started_at,
        "finished_at": finished_at,
        "summary": summary,
        "outputs": outputs,
        "selected_files": selection,
        "files": files,
        "recommendations": recommendations or [],
    }

    from forcastl.reporting.failure_metrics import (
        compute_failure_adjusted_metrics,
        records_from_detection_results,
    )
    manifest["failure_adjusted"] = compute_failure_adjusted_metrics(
        records_from_detection_results(results)
    )

    # Stable-ish filename includes model for single-model runs and uniqueness.
    safe_model = re.sub(r'[^\w\-.]', '_', model_id)
    manifest_path = config.OUTPUTS_DIR / f"run_manifest_{safe_model}_{run_id}.json"
    with open(manifest_path, "w", encoding="utf-8") as f:
        json.dump(manifest, f, indent=2)
    return str(manifest_path)
