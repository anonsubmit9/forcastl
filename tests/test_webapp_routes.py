"""Route-level tests for the FastAPI webapp (``webapp/app.py``).

These exercise the real HTTP surface a browser hits — the index page,
``/api/runs`` (list + detail), and the PDF/XLSX download routes — against a
throwaway ``outputs/`` fixture. A regression in CSV resolution, manifest
wiring, or the download builders fails here instead of only showing up as a
blank download in the browser (the exact class of bug these tests were added
to guard).

Skipped when the web deps aren't installed (``requirements-web.txt`` +
``httpx``). CI installs them, so these run in CI; the PDF case additionally
needs Chromium and self-skips when absent.
"""
import csv
import io
import json

import pytest

try:
    from fastapi.testclient import TestClient
    import forcastl.webapp.app as webapp_app
    _WEB_AVAILABLE = True
    _WEB_IMPORT_ERR = None
except Exception as e:  # pragma: no cover - exercised only without web deps
    _WEB_AVAILABLE = False
    _WEB_IMPORT_ERR = e

from forcastl.reporting.pdf_report import pdf_export_available  # core dep, safe to import

pytestmark = pytest.mark.skipif(
    not _WEB_AVAILABLE, reason=f"web deps not installed: {_WEB_IMPORT_ERR}"
)

_CSV_HEADER = [
    "Model", "Filename", "Difficulty", "Test Set", "Status", "Malicious",
    "Confidence %", "Alignment %", "Structure Valid", "Event IDs", "Processes",
    "Accounts", "Commands", "Network", "Registry", "Hallucination %",
    "Response Time (s)", "Extraction (6)", "Interpretation (6)",
    "NoHallucination (4)", "Reasoning (4)", "Total (20)", "Grade", "LLM Response",
]


def _write_csv(path):
    with open(path, "w", newline="", encoding="utf-8") as f:
        f.write("# detection results - glossary line skipped by read_detection_csv\n")
        w = csv.writer(f)
        w.writerow(_CSV_HEADER)
        w.writerow(["testmodel", "TA0006/attack1.csv", "easy", "A", "processed", "YES",
                    "90.0", "100.0", "YES", "4624", "lsass.exe", "admin", "whoami", "", "",
                    "0.0", "1.20", "6", "6", "4", "4", "20", "A", "MALICIOUS: YES"])
        w.writerow(["testmodel", "_benign/BENIGN-x.csv", "easy", "D", "processed", "NO",
                    "80.0", "100.0", "YES", "4624", "", "user", "", "", "",
                    "0.0", "0.90", "6", "6", "4", "4", "20", "A", "MALICIOUS: NO"])


def _manifest(run_id, csv_path):
    return {
        "schema_version": 1, "run_id": run_id, "kind": "detection", "raw_only": True,
        "server": "http://localhost:1234", "model": "testmodel", "input_format": "csv",
        "csv_dir": "data/csv", "mode": "sample", "sample_size": 2,
        "max_output_tokens": 1000, "timeout_s": 300, "delay_s": 0.2,
        "difficulty_filter": None, "test_set_filter": None,
        "started_at": "2026-01-01T00:00:00", "finished_at": "2026-01-01T00:01:00",
        "summary": {
            "total_files": 2, "processed": 2, "error": 0, "malicious": 1, "benign": 1,
            "avg_alignment": 100.0, "avg_hallucination_rate": 0.0, "avg_score_20": 20.0,
            "grade_dist": {"A": 2}, "recall_pct": 100.0, "fp_rate_pct": 0.0,
            "precision_pct": 100.0, "total_malicious_in_scope": 1,
            "verdict": {"tier": "insufficient", "insufficient": True, "reasons": []},
        },
        "outputs": ({"csv_results": str(csv_path)} if csv_path else {}),
        "selected_files": ["attack1.csv", "BENIGN-x.csv"], "files": [],
        "recommendations": [], "failure_adjusted": {},
    }


@pytest.fixture
def client(tmp_path, monkeypatch):
    """A TestClient wired to a temp outputs/ with one full run and one CSV-less run."""
    outputs = (tmp_path / "outputs")
    outputs.mkdir()

    csv_path = outputs / "detection_results_testmodel_20260101_000000.csv"
    _write_csv(csv_path)
    (outputs / "run_manifest_testmodel_20260101_000000.json").write_text(
        json.dumps(_manifest("20260101_000000", csv_path)), encoding="utf-8")
    # A second run with no CSV recorded, to exercise the xlsx 404 path.
    (outputs / "run_manifest_testmodel_20260101_010000.json").write_text(
        json.dumps(_manifest("20260101_010000", None)), encoding="utf-8")

    # Ground truth is keyed by .evtx even for CSV inputs — fixture mirrors that
    # so the per-file appendix's .csv->.evtx fallback is exercised by /pdf.
    gt_path = tmp_path / "ground_truth_evidence.json"
    gt_path.write_text(json.dumps({"files": {
        "attack1.evtx": {"malicious": "YES"},
        "BENIGN-x.evtx": {"malicious": "NO"},
    }}), encoding="utf-8")

    monkeypatch.setattr(webapp_app, "OUTPUTS_DIR", outputs.resolve())
    monkeypatch.setattr(webapp_app.config, "GROUND_TRUTH_FILE", str(gt_path))
    webapp_app._invalidate_run_id_index()
    try:
        yield TestClient(webapp_app.app)
    finally:
        webapp_app._invalidate_run_id_index()


# ── index + listing ────────────────────────────────────────────────────────

def test_index_serves_page(client):
    r = client.get("/")
    assert r.status_code == 200
    assert "FORCAST-L" in r.text


def test_list_runs(client):
    r = client.get("/api/runs")
    assert r.status_code == 200
    ids = {run["run_id"] for run in r.json()}
    assert {"20260101_000000", "20260101_010000"} <= ids


def test_run_detail(client):
    r = client.get("/api/runs/20260101_000000")
    assert r.status_code == 200
    data = r.json()
    assert data["model"] == "testmodel"
    assert data["summary"]["verdict"]["tier"] == "insufficient"
    assert data["summary"]["recall_pct"] == 100.0


def test_run_detail_unknown_404(client):
    assert client.get("/api/runs/nope-not-a-run").status_code == 404


# ── XLSX download (CI-runnable: no Chromium needed) ─────────────────────────

def test_xlsx_download_is_valid_workbook(client):
    r = client.get("/api/runs/run_manifest_testmodel_20260101_000000.json/xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Request", "Results"]
    values = [c.value for row in wb["Results"].iter_rows() for c in row]
    assert "TA0006/attack1.csv" in values  # real per-file row made it into the sheet


def test_xlsx_404_when_run_has_no_csv(client):
    r = client.get("/api/runs/run_manifest_testmodel_20260101_010000.json/xlsx")
    assert r.status_code == 404


# ── traversal / input guards ────────────────────────────────────────────────

def test_pdf_route_rejects_dotdot(client):
    # The handler rejects any manifest name containing ".." before touching disk.
    assert client.get("/api/runs/a..b.json/pdf").status_code == 400


def test_pdf_route_unknown_manifest_404(client):
    assert client.get("/api/runs/run_manifest_missing_20990101_000000.json/pdf").status_code == 404


# ── PDF download (needs Chromium; self-skips in CI) ─────────────────────────

@pytest.mark.skipif(not pdf_export_available(), reason="Chromium not installed")
def test_pdf_download_renders(client):
    r = client.get("/api/runs/run_manifest_testmodel_20260101_000000.json/pdf")
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    assert r.content[:5] == b"%PDF-"
    assert len(r.content) > 1000
